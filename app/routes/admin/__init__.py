from flask import Blueprint, render_template, request, jsonify, redirect, url_for, flash, session, current_app
from functools import wraps
from app.extensions import db
from app.models.user import AdminUser

admin_bp = Blueprint('admin', __name__)


def admin_login_required(f):
    """Decorator: requires admin to be logged in via session."""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get('admin_logged_in'):
            return redirect(url_for('admin.login'))
        return f(*args, **kwargs)
    return decorated_function


def get_current_admin():
    """Get the currently logged-in admin user."""
    admin_id = session.get('admin_id')
    if admin_id:
        return AdminUser.query.get(admin_id)
    return None


@admin_bp.context_processor
def inject_admin():
    """Inject admin user into all admin templates."""
    return {'current_admin': get_current_admin()}


# ── Admin Login / Logout ────────────────────────────────────────────────────

@admin_bp.route('/login', methods=['GET', 'POST'])
def login():
    if session.get('admin_logged_in'):
        return redirect(url_for('admin.dashboard'))

    if request.method == 'POST':
        email = request.form.get('email', '').strip().lower()
        password = request.form.get('password', '')

        admin = AdminUser.query.filter_by(email=email, is_active=True).first()
        if admin and admin.check_password(password):
            session['admin_logged_in'] = True
            session['admin_id'] = admin.id
            session['admin_name'] = admin.name
            session['admin_role'] = admin.role
            session.permanent = True
            flash('Welcome back, ' + admin.name, 'success')
            return redirect(url_for('admin.dashboard'))
        else:
            flash('Invalid email or password.', 'danger')

    return render_template('admin/login.html')


@admin_bp.route('/logout')
def logout():
    session.pop('admin_logged_in', None)
    session.pop('admin_id', None)
    session.pop('admin_name', None)
    session.pop('admin_role', None)
    flash('You have been logged out from the admin panel.', 'info')
    return redirect(url_for('admin.login'))


# ── Admin Dashboard ─────────────────────────────────────────────────────────

@admin_bp.route('/')
@admin_bp.route('/dashboard')
@admin_login_required
def dashboard():
    from app.models.guru import Guru
    from app.models.lekhana import LekhanaSession, SevaProgress
    from app.models.user import User

    stats = {
        'total_users': User.query.count(),
        'total_gurus': Guru.query.count(),
        'verified_gurus': Guru.query.filter_by(is_verified=True).count(),
        'total_sessions': LekhanaSession.query.count(),
        'completed_sessions': LekhanaSession.query.filter_by(status='completed').count(),
    }

    recent_sessions = (LekhanaSession.query
                       .order_by(LekhanaSession.created_at.desc())
                       .limit(10).all())

    return render_template('admin/dashboard.html', stats=stats, recent_sessions=recent_sessions)


# ── Guru Management ─────────────────────────────────────────────────────────

@admin_bp.route('/gurus')
@admin_login_required
def gurus_list():
    from app.models.guru import Guru
    gurus = Guru.query.order_by(Guru.guru_order.asc()).all()
    return render_template('admin/gurus/list.html', gurus=gurus)


@admin_bp.route('/gurus/add', methods=['GET', 'POST'])
@admin_login_required
def guru_add():
    from app.models.guru import Guru, GuruWork, GuruSource
    from slugify import slugify

    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        if not name:
            flash('Guru name is required.', 'danger')
            return redirect(url_for('admin.guru_add'))

        slug = slugify(name)
        # Ensure unique slug
        if Guru.query.filter_by(slug=slug).first():
            slug = slug + '-' + str(Guru.query.count() + 1)

        guru = Guru(
            name=name,
            traditional_name=request.form.get('traditional_name', '').strip() or None,
            slug=slug,
            guru_order=int(request.form.get('guru_order') or 0),
            short_description=request.form.get('short_description', '').strip() or None,
            biography=request.form.get('biography', '').strip() or None,
            birth_date=request.form.get('birth_date', '').strip() or None,
            aradhana_date=request.form.get('aradhana_date', '').strip() or None,
            lekhana_text=request.form.get('lekhana_text', '').strip() or None,
            lekhana_sanskrit=request.form.get('lekhana_sanskrit', '').strip() or None,
            lekhana_kannada=request.form.get('lekhana_kannada', '').strip() or None,
            lekhana_english=request.form.get('lekhana_english', '').strip() or None,
            is_verified=bool(request.form.get('is_verified')),
            is_active=bool(request.form.get('is_active', True)),
        )
        db.session.add(guru)
        db.session.commit()
        _log_action('create', 'gurus', guru.id, f'Created guru: {guru.name}')
        flash(f'Guru "{guru.name}" created successfully.', 'success')
        return redirect(url_for('admin.gurus_list'))

    gurus_all = Guru.query.order_by(Guru.guru_order.asc()).all()
    return render_template('admin/gurus/form.html', guru=None, gurus_all=gurus_all, action='add')


@admin_bp.route('/gurus/<int:guru_id>/edit', methods=['GET', 'POST'])
@admin_login_required
def guru_edit(guru_id):
    from app.models.guru import Guru

    guru = Guru.query.get_or_404(guru_id)

    if request.method == 'POST':
        guru.name = request.form.get('name', '').strip() or guru.name
        guru.traditional_name = request.form.get('traditional_name', '').strip() or None
        guru.guru_order = int(request.form.get('guru_order') or guru.guru_order)
        guru.short_description = request.form.get('short_description', '').strip() or None
        guru.biography = request.form.get('biography', '').strip() or None
        guru.birth_date = request.form.get('birth_date', '').strip() or None
        guru.aradhana_date = request.form.get('aradhana_date', '').strip() or None
        guru.lekhana_text = request.form.get('lekhana_text', '').strip() or None
        guru.lekhana_sanskrit = request.form.get('lekhana_sanskrit', '').strip() or None
        guru.lekhana_kannada = request.form.get('lekhana_kannada', '').strip() or None
        guru.lekhana_english = request.form.get('lekhana_english', '').strip() or None
        guru.is_verified = bool(request.form.get('is_verified'))
        guru.is_active = bool(request.form.get('is_active', True))

        prev_id = request.form.get('previous_guru_id')
        next_id = request.form.get('next_guru_id')
        guru.previous_guru_id = int(prev_id) if prev_id else None
        guru.next_guru_id = int(next_id) if next_id else None

        db.session.commit()
        _log_action('update', 'gurus', guru.id, f'Updated guru: {guru.name}')
        flash(f'Guru "{guru.name}" updated successfully.', 'success')
        return redirect(url_for('admin.gurus_list'))

    gurus_all = Guru.query.filter(Guru.id != guru_id).order_by(Guru.guru_order.asc()).all()
    return render_template('admin/gurus/form.html', guru=guru, gurus_all=gurus_all, action='edit')


@admin_bp.route('/gurus/<int:guru_id>/delete', methods=['POST'])
@admin_login_required
def guru_delete(guru_id):
    from app.models.guru import Guru
    guru = Guru.query.get_or_404(guru_id)
    name = guru.name
    db.session.delete(guru)
    db.session.commit()
    _log_action('delete', 'gurus', guru_id, f'Deleted guru: {name}')
    flash(f'Guru "{name}" deleted.', 'warning')
    return redirect(url_for('admin.gurus_list'))


@admin_bp.route('/gurus/<int:guru_id>/verify', methods=['POST'])
@admin_login_required
def guru_verify(guru_id):
    from app.models.guru import Guru
    guru = Guru.query.get_or_404(guru_id)
    guru.is_verified = not guru.is_verified
    db.session.commit()
    status = 'verified' if guru.is_verified else 'unverified'
    _log_action('verify', 'gurus', guru_id, f'Guru {status}: {guru.name}')
    flash(f'Guru "{guru.name}" is now {status}.', 'success')
    return redirect(url_for('admin.gurus_list'))


@admin_bp.route('/gurus/<int:guru_id>/image', methods=['POST'])
@admin_login_required
def guru_image_upload(guru_id):
    import os
    from werkzeug.utils import secure_filename
    from app.models.guru import Guru

    guru = Guru.query.get_or_404(guru_id)

    if 'image' not in request.files:
        flash('No image file provided.', 'danger')
        return redirect(url_for('admin.guru_edit', guru_id=guru_id))

    file = request.files['image']
    if file.filename == '':
        flash('No file selected.', 'danger')
        return redirect(url_for('admin.guru_edit', guru_id=guru_id))

    allowed = {'png', 'jpg', 'jpeg', 'gif', 'webp'}
    ext = file.filename.rsplit('.', 1)[-1].lower() if '.' in file.filename else ''
    if ext not in allowed:
        flash('Invalid file type. Use PNG, JPG, GIF, or WebP.', 'danger')
        return redirect(url_for('admin.guru_edit', guru_id=guru_id))

    filename = secure_filename(f"guru_{guru_id}_{guru.slug}.{ext}")
    upload_path = os.path.join(current_app.config.get('UPLOAD_FOLDER', 'uploads'), 'gurus')
    os.makedirs(upload_path, exist_ok=True)
    file.save(os.path.join(upload_path, filename))

    guru.image_url = f"/uploads/gurus/{filename}"
    db.session.commit()
    flash('Image uploaded successfully.', 'success')
    return redirect(url_for('admin.guru_edit', guru_id=guru_id))


# ── Content Management ──────────────────────────────────────────────────────

@admin_bp.route('/content/books')
@admin_login_required
def books_list():
    from app.models.content import Book
    books = Book.query.order_by(Book.created_at.desc()).all()
    return render_template('admin/content/books.html', books=books)


@admin_bp.route('/content/books/add', methods=['GET', 'POST'])
@admin_login_required
def book_add():
    from app.models.content import Book
    if request.method == 'POST':
        book = Book(
            title=request.form.get('title', '').strip(),
            author=request.form.get('author', '').strip() or None,
            category=request.form.get('category', '').strip() or None,
            language=request.form.get('language', 'en'),
            description=request.form.get('description', '').strip() or None,
            file_url=request.form.get('file_url', '').strip() or None,
            is_active=bool(request.form.get('is_active', True)),
        )
        db.session.add(book)
        db.session.commit()
        flash('Book added successfully.', 'success')
        return redirect(url_for('admin.books_list'))
    return render_template('admin/content/book_form.html', book=None)


@admin_bp.route('/content/books/<int:book_id>/edit', methods=['GET', 'POST'])
@admin_login_required
def book_edit(book_id):
    from app.models.content import Book
    book = Book.query.get_or_404(book_id)
    if request.method == 'POST':
        book.title = request.form.get('title', book.title)
        book.author = request.form.get('author', '').strip() or None
        book.category = request.form.get('category', '').strip() or None
        book.language = request.form.get('language', 'en')
        book.description = request.form.get('description', '').strip() or None
        book.file_url = request.form.get('file_url', '').strip() or None
        book.is_active = bool(request.form.get('is_active'))
        db.session.commit()
        flash('Book updated.', 'success')
        return redirect(url_for('admin.books_list'))
    return render_template('admin/content/book_form.html', book=book)


@admin_bp.route('/content/books/<int:book_id>/delete', methods=['POST'])
@admin_login_required
def book_delete(book_id):
    from app.models.content import Book
    book = Book.query.get_or_404(book_id)
    db.session.delete(book)
    db.session.commit()
    flash('Book deleted.', 'warning')
    return redirect(url_for('admin.books_list'))


@admin_bp.route('/content/media')
@admin_login_required
def media_list():
    from app.models.content import Media
    media = Media.query.order_by(Media.created_at.desc()).all()
    return render_template('admin/content/media.html', media_list=media)


@admin_bp.route('/content/media/add', methods=['GET', 'POST'])
@admin_login_required
def media_add():
    from app.models.content import Media
    if request.method == 'POST':
        item = Media(
            title=request.form.get('title', '').strip(),
            speaker=request.form.get('speaker', '').strip() or None,
            topic=request.form.get('topic', '').strip() or None,
            language=request.form.get('language', 'en'),
            media_type=request.form.get('media_type', 'video'),
            embed_url=request.form.get('embed_url', '').strip() or None,
            thumbnail_url=request.form.get('thumbnail_url', '').strip() or None,
            duration=request.form.get('duration', '').strip() or None,
            is_active=bool(request.form.get('is_active', True)),
        )
        db.session.add(item)
        db.session.commit()
        flash('Media item added successfully.', 'success')
        return redirect(url_for('admin.media_list'))
    return render_template('admin/content/media_form.html', media=None)


@admin_bp.route('/content/events')
@admin_login_required
def events_list():
    from app.models.content import CalendarEvent
    events = CalendarEvent.query.order_by(CalendarEvent.event_date.asc()).all()
    return render_template('admin/content/events.html', events=events)


@admin_bp.route('/content/events/add', methods=['GET', 'POST'])
@admin_login_required
def event_add():
    from app.models.content import CalendarEvent
    from app.models.guru import Guru
    if request.method == 'POST':
        event = CalendarEvent(
            title=request.form.get('title', '').strip(),
            event_date=request.form.get('event_date'),
            tithi=request.form.get('tithi', '').strip() or None,
            event_type=request.form.get('event_type', 'event'),
            guru_id=int(request.form.get('guru_id')) if request.form.get('guru_id') else None,
            description=request.form.get('description', '').strip() or None,
            is_active=bool(request.form.get('is_active', True)),
        )
        db.session.add(event)
        db.session.commit()
        flash('Event added successfully.', 'success')
        return redirect(url_for('admin.events_list'))
    gurus = Guru.query.filter_by(is_verified=True).order_by(Guru.guru_order).all()
    return render_template('admin/content/event_form.html', event=None, gurus=gurus)


@admin_bp.route('/content/events/<int:event_id>/delete', methods=['POST'])
@admin_login_required
def event_delete(event_id):
    from app.models.content import CalendarEvent
    event = CalendarEvent.query.get_or_404(event_id)
    db.session.delete(event)
    db.session.commit()
    flash('Event deleted.', 'warning')
    return redirect(url_for('admin.events_list'))


# ── User Management ─────────────────────────────────────────────────────────

@admin_bp.route('/users')
@admin_login_required
def users_list():
    from app.models.user import User
    page = request.args.get('page', 1, type=int)
    users = User.query.order_by(User.created_at.desc()).paginate(page=page, per_page=20)
    return render_template('admin/users/list.html', users=users)


@admin_bp.route('/users/<int:user_id>/toggle', methods=['POST'])
@admin_login_required
def user_toggle(user_id):
    from app.models.user import User
    user = User.query.get_or_404(user_id)
    user.is_active = not user.is_active
    db.session.commit()
    status = 'activated' if user.is_active else 'deactivated'
    flash(f'User "{user.name}" {status}.', 'success')
    return redirect(url_for('admin.users_list'))


# ── Analytics ────────────────────────────────────────────────────────────────

@admin_bp.route('/analytics')
@admin_login_required
def analytics():
    from app.models.user import User
    from app.models.guru import Guru
    from app.models.lekhana import LekhanaSession
    from sqlalchemy import func

    stats = {
        'total_users': User.query.count(),
        'active_users': User.query.filter_by(is_active=True).count(),
        'total_gurus': Guru.query.count(),
        'verified_gurus': Guru.query.filter_by(is_verified=True).count(),
        'total_sessions': LekhanaSession.query.count(),
        'completed_sessions': LekhanaSession.query.filter_by(status='completed').count(),
        'in_progress_sessions': LekhanaSession.query.filter_by(status='in_progress').count(),
    }

    # Top gurus by session count
    top_gurus = (db.session.query(Guru.name, func.count(LekhanaSession.id).label('session_count'))
                 .join(LekhanaSession, Guru.id == LekhanaSession.guru_id)
                 .group_by(Guru.id)
                 .order_by(func.count(LekhanaSession.id).desc())
                 .limit(10).all())

    return render_template('admin/analytics.html', stats=stats, top_gurus=top_gurus)


# ── AI Knowledge Base ────────────────────────────────────────────────────────

@admin_bp.route('/ai-knowledge-base')
@admin_login_required
def ai_kb():
    from app.models.ai import AIDocument
    docs = AIDocument.query.order_by(AIDocument.created_at.desc()).all()
    return render_template('admin/ai_kb.html', docs=docs)


@admin_bp.route('/ai-knowledge-base/upload', methods=['POST'])
@admin_login_required
def ai_kb_upload():
    import os
    from werkzeug.utils import secure_filename
    from app.models.ai import AIDocument, AIChunk
    from app.services.ai_service import get_ai_service

    if 'file' not in request.files:
        flash('No file provided.', 'danger')
        return redirect(url_for('admin.ai_kb'))

    file = request.files['file']
    if file.filename == '':
        flash('No file selected.', 'danger')
        return redirect(url_for('admin.ai_kb'))

    allowed = {'pdf', 'txt', 'md', 'docx'}
    ext = file.filename.rsplit('.', 1)[-1].lower() if '.' in file.filename else ''
    if ext not in allowed:
        flash('Only PDF, TXT, MD, and DOCX files are allowed.', 'danger')
        return redirect(url_for('admin.ai_kb'))

    filename = secure_filename(file.filename)
    upload_path = os.path.join(current_app.config.get('UPLOAD_FOLDER', 'uploads'), 'ai_docs')
    os.makedirs(upload_path, exist_ok=True)
    file_path = os.path.join(upload_path, filename)
    file.save(file_path)

    # Create DB record
    title = request.form.get('title', filename).strip()
    doc = AIDocument(title=title, source_type=ext, file_path=file_path, is_indexed=False)
    db.session.add(doc)
    db.session.commit()

    # Extract text and index
    text = _extract_text(file_path, ext)
    if text:
        chunks = _chunk_text(text)
        chunk_dicts = [{'text': c, 'index': i} for i, c in enumerate(chunks)]

        ai_svc = get_ai_service()
        success = ai_svc.add_document(str(doc.id), doc.title, chunk_dicts)

        if success:
            # Save chunks to DB
            for i, chunk_text in enumerate(chunks):
                chunk = AIChunk(
                    document_id=doc.id,
                    chunk_text=chunk_text,
                    chunk_index=i,
                    embedding_id=f"doc_{doc.id}_chunk_{i}"
                )
                db.session.add(chunk)
            doc.is_indexed = True
            db.session.commit()
            flash(f'Document "{title}" uploaded and indexed ({len(chunks)} chunks).', 'success')
        else:
            flash(f'Document uploaded but indexing failed. AI service may not be configured.', 'warning')
    else:
        flash('Document uploaded but text extraction failed.', 'warning')

    return redirect(url_for('admin.ai_kb'))


@admin_bp.route('/ai-knowledge-base/<int:doc_id>/delete', methods=['POST'])
@admin_login_required
def ai_kb_delete(doc_id):
    from app.models.ai import AIDocument, AIChunk
    doc = AIDocument.query.get_or_404(doc_id)
    AIChunk.query.filter_by(document_id=doc_id).delete()
    db.session.delete(doc)
    db.session.commit()
    flash('Document removed from knowledge base.', 'warning')
    return redirect(url_for('admin.ai_kb'))


@admin_bp.route('/ai-knowledge-base/index-gurus', methods=['POST'])
@admin_login_required
def ai_index_gurus():
    """Trigger re-indexing of all verified Guru data."""
    from app.services.ai_service import get_ai_service
    ai_svc = get_ai_service()
    result = ai_svc.index_guru_data()
    if result.get('success'):
        flash(f'Successfully indexed {result["indexed"]} Guru entries.', 'success')
    else:
        flash(f'Indexing failed: {result.get("message", "Unknown error")}', 'danger')
    return redirect(url_for('admin.ai_kb'))


# ── Helper functions ─────────────────────────────────────────────────────────

def _log_action(action, table, target_id, details=''):
    """Log an admin action to audit_logs."""
    try:
        from app.models.ai import AuditLog
        log = AuditLog(
            admin_id=session.get('admin_id'),
            action=action,
            target_table=table,
            target_id=target_id,
            details=details
        )
        db.session.add(log)
        db.session.commit()
    except Exception:
        pass  # Don't fail the main action if logging fails


def _extract_text(file_path: str, ext: str) -> str:
    """Extract text from uploaded document."""
    try:
        if ext == 'txt' or ext == 'md':
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        elif ext == 'pdf':
            try:
                import PyPDF2
                text = ''
                with open(file_path, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    for page in reader.pages:
                        text += page.extract_text() or ''
                return text
            except ImportError:
                # Try pdfplumber as fallback
                try:
                    import pdfplumber
                    text = ''
                    with pdfplumber.open(file_path) as pdf:
                        for page in pdf.pages:
                            text += page.extract_text() or ''
                    return text
                except ImportError:
                    return ''
        elif ext == 'docx':
            try:
                import docx
                doc = docx.Document(file_path)
                return '\n'.join(p.text for p in doc.paragraphs)
            except ImportError:
                return ''
    except Exception:
        return ''
    return ''


def _chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> list:
    """Split text into overlapping chunks for vector indexing."""
    words = text.split()
    chunks = []
    for i in range(0, len(words), chunk_size - overlap):
        chunk = ' '.join(words[i:i + chunk_size])
        if chunk.strip():
            chunks.append(chunk)
    return chunks
